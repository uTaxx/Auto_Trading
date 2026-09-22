"""백테스트 결과를 서식 있는 엑셀로 만든다.

화면에 남는 JSON과 별개로, 구글 드라이브에 보관해 두고 나중에 다시
열어 보기 좋은 형태로 남기는 것이 목적이다. 헤더에 배경색을 주고,
금액에는 천단위 구분을, 수익률에는 퍼센트 서식을 입힌다.
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="2F6F65")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)
POSITIVE_FONT = Font(color="1B7A43")
NEGATIVE_FONT = Font(color="B5502E")
MONEY_FORMAT = "#,##0"
SHARE_FORMAT = "#,##0.0000"
PERCENT_FORMAT = "0.00%"


def _set_column_widths(ws: Worksheet, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _write_header_row(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col_idx, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=text)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _write_summary_block(ws: Worksheet, title: str, info_lines: list[str]) -> int:
    """제목과 안내 줄을 적고, 다음으로 쓸 수 있는 줄 번호를 돌려준다."""
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    row = 2
    for line in info_lines:
        ws.cell(row=row, column=1, value=line)
        row += 1
    return row + 1


def build_symbol_report(
    symbol: str,
    strategy_name: str,
    capital: float,
    start: str,
    end: str,
    generated_at_kst: str,
    summary_row: dict,
    daily: pd.DataFrame,
) -> Workbook:
    """종목 하나, 전략 하나의 분석 결과를 담은 엑셀을 만든다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"

    table_row = _write_summary_block(
        ws,
        f"{symbol} - {strategy_name} 분석 결과",
        [f"조회기간: {start} ~ {end}", f"총자본: {capital:,.0f}원", f"생성시각(KST): {generated_at_kst}"],
    )
    headers = [
        "종목", "매수방식", "매수금액", "매수빈도", "이동평균조건", "등락구간", "익절선", "손절선",
        "총투자금", "실현손익", "평가손익", "합계", "수익률", "최대낙폭", "현금부족일수",
    ]
    _write_header_row(ws, table_row, headers)
    r = table_row + 1
    ws.cell(row=r, column=1, value=summary_row["symbol"])
    ws.cell(row=r, column=2, value=summary_row.get("매수방식"))
    if summary_row.get("매수금액") is not None:
        ws.cell(row=r, column=3, value=summary_row["매수금액"]).number_format = MONEY_FORMAT
    if summary_row.get("매수빈도") is not None:
        ws.cell(row=r, column=4, value=f"{summary_row['매수빈도']}일")
    ws.cell(row=r, column=5, value=summary_row.get("이동평균조건"))
    ws.cell(row=r, column=6, value=summary_row.get("등락구간"))
    ws.cell(row=r, column=7, value=summary_row.get("익절선"))
    ws.cell(row=r, column=8, value=summary_row.get("손절선"))
    ws.cell(row=r, column=9, value=summary_row["총투자금"]).number_format = MONEY_FORMAT
    ws.cell(row=r, column=10, value=summary_row["실현손익"]).number_format = MONEY_FORMAT
    ws.cell(row=r, column=11, value=summary_row["평가손익"]).number_format = MONEY_FORMAT
    ws.cell(row=r, column=12, value=summary_row["합계"]).number_format = MONEY_FORMAT
    ws.cell(row=r, column=13, value=summary_row["수익률"] / 100).number_format = PERCENT_FORMAT
    ws.cell(row=r, column=14, value=summary_row.get("최대낙폭", 0) / 100).number_format = PERCENT_FORMAT
    ws.cell(row=r, column=15, value=summary_row.get("현금부족일수", 0))
    _set_column_widths(ws, [10, 18, 12, 10, 22, 22, 9, 9, 14, 14, 14, 14, 10, 10, 12])

    ws2 = wb.create_sheet("일별 시계열")
    daily_headers = [
        "거래일", "종가", "당일매수금액", "당일매수부족금액", "당일매도금액", "매도유형", "현금", "보유수량",
        "평균단가", "누적투자금", "평가금액", "실현손익", "평가손익", "손익합계", "총자산", "누적수익률",
    ]
    _write_header_row(ws2, 1, daily_headers)
    for r_idx, row_data in enumerate(daily.itertuples(index=False), start=2):
        buy_amount = float(getattr(row_data, "buy_amount", 0.0) or 0.0)
        buy_shortfall = float(getattr(row_data, "buy_shortfall", 0.0) or 0.0)
        sell_amount = float(getattr(row_data, "sell_amount", 0.0) or 0.0)
        sell_type = getattr(row_data, "sell_type", None)
        if pd.isna(sell_type):  # 매도 없는 날은 pandas가 None을 NaN으로 바꿔 둔다
            sell_type = None
        total_pnl = float(row_data.total_pnl)

        ws2.cell(row=r_idx, column=1, value=str(row_data.trade_date))
        ws2.cell(row=r_idx, column=2, value=round(float(row_data.close))).number_format = MONEY_FORMAT
        if buy_amount > 0:
            ws2.cell(row=r_idx, column=3, value=round(buy_amount)).number_format = MONEY_FORMAT
        if buy_shortfall > 0:
            cell = ws2.cell(row=r_idx, column=4, value=round(buy_shortfall))
            cell.number_format = MONEY_FORMAT
            cell.font = NEGATIVE_FONT
        if sell_amount > 0:
            ws2.cell(row=r_idx, column=5, value=round(sell_amount)).number_format = MONEY_FORMAT
        ws2.cell(row=r_idx, column=6, value=sell_type)
        ws2.cell(row=r_idx, column=7, value=round(float(row_data.cash))).number_format = MONEY_FORMAT
        ws2.cell(row=r_idx, column=8, value=float(row_data.shares)).number_format = SHARE_FORMAT
        ws2.cell(row=r_idx, column=9, value=round(float(row_data.avg_cost))).number_format = MONEY_FORMAT
        ws2.cell(row=r_idx, column=10, value=round(float(row_data.invested_cumulative))).number_format = MONEY_FORMAT
        ws2.cell(row=r_idx, column=11, value=round(float(row_data.market_value))).number_format = MONEY_FORMAT
        ws2.cell(row=r_idx, column=12, value=round(float(row_data.realized_pnl))).number_format = MONEY_FORMAT
        ws2.cell(row=r_idx, column=13, value=round(float(row_data.unrealized_pnl))).number_format = MONEY_FORMAT
        ws2.cell(row=r_idx, column=14, value=round(total_pnl)).number_format = MONEY_FORMAT
        ws2.cell(row=r_idx, column=15, value=round(float(row_data.total_value))).number_format = MONEY_FORMAT
        if capital:
            ws2.cell(row=r_idx, column=16, value=total_pnl / capital).number_format = PERCENT_FORMAT
    ws2.freeze_panes = "A2"
    _set_column_widths(ws2, [12, 12, 13, 15, 13, 9, 12, 12, 12, 14, 14, 13, 13, 13, 14, 12])

    return wb


def build_comparison_report(
    symbols: list[str],
    capital: float,
    start: str,
    end: str,
    generated_at_kst: str,
    summary_rows: list[dict],
) -> Workbook:
    """여러 종목·여러 전략을 한 번에 비교한 결과를 담은 엑셀을 만든다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "종목 비교"

    table_row = _write_summary_block(
        ws,
        f"종목 비교 결과 ({', '.join(symbols)})",
        [f"조회기간: {start} ~ {end}", f"총자본(종목·전략마다): {capital:,.0f}원", f"생성시각(KST): {generated_at_kst}"],
    )
    headers = [
        "종목", "매수방식", "매수금액", "매수빈도", "이동평균조건", "등락구간", "익절선", "손절선",
        "총투자금", "실현손익", "평가손익", "합계", "수익률", "최대낙폭", "현금부족일수",
    ]
    _write_header_row(ws, table_row, headers)
    for i, row in enumerate(summary_rows, start=table_row + 1):
        ws.cell(row=i, column=1, value=row["symbol"])
        ws.cell(row=i, column=2, value=row.get("매수방식"))
        if row.get("매수금액") is not None:
            ws.cell(row=i, column=3, value=row["매수금액"]).number_format = MONEY_FORMAT
        if row.get("매수빈도") is not None:
            ws.cell(row=i, column=4, value=f"{row['매수빈도']}일")
        ws.cell(row=i, column=5, value=row.get("이동평균조건"))
        ws.cell(row=i, column=6, value=row.get("등락구간"))
        ws.cell(row=i, column=7, value=row.get("익절선"))
        ws.cell(row=i, column=8, value=row.get("손절선"))
        ws.cell(row=i, column=9, value=row["총투자금"]).number_format = MONEY_FORMAT
        ws.cell(row=i, column=10, value=row["실현손익"]).number_format = MONEY_FORMAT
        ws.cell(row=i, column=11, value=row["평가손익"]).number_format = MONEY_FORMAT
        ws.cell(row=i, column=12, value=row["합계"]).number_format = MONEY_FORMAT
        pct_cell = ws.cell(row=i, column=13, value=row["수익률"] / 100)
        pct_cell.number_format = PERCENT_FORMAT
        pct_cell.font = POSITIVE_FONT if row["수익률"] >= 0 else NEGATIVE_FONT
        ws.cell(row=i, column=14, value=row.get("최대낙폭", 0) / 100).number_format = PERCENT_FORMAT
        shortfall_days_cell = ws.cell(row=i, column=15, value=row.get("현금부족일수", 0))
        if row.get("현금부족일수", 0) > 0:
            shortfall_days_cell.font = NEGATIVE_FONT
    ws.freeze_panes = f"A{table_row + 1}"
    _set_column_widths(ws, [10, 18, 12, 10, 22, 22, 9, 9, 14, 14, 14, 14, 10, 10, 12])

    return wb


def build_portfolio_report(
    symbols: list[str],
    capital: float,
    start: str,
    end: str,
    generated_at_kst: str,
    summary: dict,
    symbol_breakdown: dict[str, dict],
    daily: pd.DataFrame,
) -> Workbook:
    """여러 종목이 현금 하나를 나눠 쓴 포트폴리오 백테스트 결과를 담은
    엑셀을 만든다(`portfolio.run_portfolio_backtest`의 결과). 종목마다
    따로 총자본을 가정하는 `build_comparison_report`와 달리, 포트폴리오
    전체 총자본 하나와 종목별 실제 배분 결과를 같이 보여준다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "포트폴리오 요약"

    table_row = _write_summary_block(
        ws,
        f"포트폴리오 백테스트 결과 ({', '.join(symbols)})",
        [
            f"조회기간: {start} ~ {end}",
            f"총자본(공유): {capital:,.0f}원",
            f"생성시각(KST): {generated_at_kst}",
            (
                f"전체 수익률: {summary['수익률']:.2f}%, 최대낙폭: {summary.get('최대낙폭', 0):.2f}%, "
                f"현금부족일수: {summary.get('현금부족일수', 0)}일"
            ),
        ],
    )
    headers = ["종목", "총투자금", "실현손익", "평가손익", "합계"]
    _write_header_row(ws, table_row, headers)
    r = table_row + 1
    for symbol in symbols:
        row = symbol_breakdown.get(symbol, {})
        ws.cell(row=r, column=1, value=symbol)
        ws.cell(row=r, column=2, value=row.get("총투자금", 0)).number_format = MONEY_FORMAT
        ws.cell(row=r, column=3, value=row.get("실현손익", 0)).number_format = MONEY_FORMAT
        ws.cell(row=r, column=4, value=row.get("평가손익", 0)).number_format = MONEY_FORMAT
        ws.cell(row=r, column=5, value=row.get("합계", 0)).number_format = MONEY_FORMAT
        r += 1
    ws.cell(row=r, column=1, value="합계(포트폴리오)").font = Font(bold=True)
    ws.cell(row=r, column=2, value=summary["총투자금"]).number_format = MONEY_FORMAT
    ws.cell(row=r, column=3, value=summary["실현손익"]).number_format = MONEY_FORMAT
    ws.cell(row=r, column=4, value=summary["평가손익"]).number_format = MONEY_FORMAT
    ws.cell(row=r, column=5, value=summary["합계"]).number_format = MONEY_FORMAT
    _set_column_widths(ws, [12, 16, 16, 16, 16])

    ws2 = wb.create_sheet("일별 시계열")
    daily_headers = [
        "거래일", "당일매수금액", "당일매수부족금액", "당일매도금액", "현금", "평가금액",
        "누적투자금", "실현손익", "평가손익", "손익합계", "총자산", "누적수익률",
    ]
    _write_header_row(ws2, 1, daily_headers)
    for r_idx, row_data in enumerate(daily.itertuples(index=False), start=2):
        buy_amount = float(row_data.buy_amount or 0.0)
        buy_shortfall = float(row_data.buy_shortfall or 0.0)
        sell_amount = float(row_data.sell_amount or 0.0)
        total_pnl = float(row_data.total_pnl)

        ws2.cell(row=r_idx, column=1, value=str(row_data.trade_date))
        if buy_amount > 0:
            ws2.cell(row=r_idx, column=2, value=round(buy_amount)).number_format = MONEY_FORMAT
        if buy_shortfall > 0:
            cell = ws2.cell(row=r_idx, column=3, value=round(buy_shortfall))
            cell.number_format = MONEY_FORMAT
            cell.font = NEGATIVE_FONT
        if sell_amount > 0:
            ws2.cell(row=r_idx, column=4, value=round(sell_amount)).number_format = MONEY_FORMAT
        ws2.cell(row=r_idx, column=5, value=round(float(row_data.cash))).number_format = MONEY_FORMAT
        ws2.cell(row=r_idx, column=6, value=round(float(row_data.market_value))).number_format = MONEY_FORMAT
        ws2.cell(row=r_idx, column=7, value=round(float(row_data.invested_cumulative))).number_format = MONEY_FORMAT
        ws2.cell(row=r_idx, column=8, value=round(float(row_data.realized_pnl))).number_format = MONEY_FORMAT
        ws2.cell(row=r_idx, column=9, value=round(float(row_data.unrealized_pnl))).number_format = MONEY_FORMAT
        ws2.cell(row=r_idx, column=10, value=round(total_pnl)).number_format = MONEY_FORMAT
        ws2.cell(row=r_idx, column=11, value=round(float(row_data.total_value))).number_format = MONEY_FORMAT
        if capital:
            ws2.cell(row=r_idx, column=12, value=total_pnl / capital).number_format = PERCENT_FORMAT
    ws2.freeze_panes = "A2"
    _set_column_widths(ws2, [12, 14, 15, 13, 13, 13, 14, 13, 13, 13, 14, 12])

    return wb
